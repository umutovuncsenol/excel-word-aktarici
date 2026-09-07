"""Excel verilerini tarih üzerinden bir Word tablosuna aktarır.

Gerekli paketler:
    pip install openpyxl python-docx

Program ağ bağlantısı kullanmaz. Seçilen dosyalar yalnızca yerel bilgisayarda
okunur ve sonuç yeni bir DOCX dosyasına kaydedilir.
"""

from __future__ import annotations

import locale
import os
import re
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from numbers import Number
from pathlib import Path
from typing import Any

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from docx import Document
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
except ImportError as exc:
    missing_package = str(exc)
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror(
        "Eksik paket",
        "Programın çalışması için openpyxl ve python-docx gerekiyor.\n\n"
        "Komut: pip install openpyxl python-docx\n\n"
        f"Teknik ayrıntı: {missing_package}",
    )
    root.destroy()
    raise SystemExit(1) from exc


REQUIRED_HEADERS = ("Date", "PHF (€)", "Volume", "PHF VT")
VALUE_HEADERS = REQUIRED_HEADERS[1:]
MAX_HEADER_SCAN_ROWS = 100
MAX_HEADER_SCAN_COLUMNS = 100
MAX_WORD_HEADER_SCAN_ROWS = 10
OVERWRITE_EXISTING_CELLS = False


class TransferError(Exception):
    """Kullanıcıya gösterilebilecek, beklenen işlem hatası."""


@dataclass
class ExcelRecord:
    values: dict[str, str]
    blank_headers: list[str] = field(default_factory=list)


@dataclass
class ExcelData:
    sheet_name: str
    header_row: int
    records: dict[date, ExcelRecord]


@dataclass
class TransferResult:
    excel_sheet: str
    excel_header_row: int
    word_table_number: int
    word_header_row: int
    matched_dates: set[date] = field(default_factory=set)
    word_dates_without_excel: set[date] = field(default_factory=set)
    excel_dates_not_in_word: set[date] = field(default_factory=set)
    skipped_nonempty_cells: int = 0
    blank_excel_cells: int = 0
    ignored_word_rows: int = 0


def normalize_header(value: Any) -> str:
    """Başlıklardaki görünmez boşluk ve küçük yazım farklarını giderir."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text).strip().casefold()
    return text


NORMALIZED_HEADERS = {normalize_header(header): header for header in REQUIRED_HEADERS}


def parse_date(value: Any, excel_epoch: Any = None) -> date:
    """Excel veya Word'den gelen günlük tarihi standart date değerine çevirir."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    if isinstance(value, Number) and not isinstance(value, bool):
        number = float(value)
        integer = int(number)
        # 20260907 biçimindeki sayısal tarihler.
        if 19000101 <= integer <= 29991231 and number == integer:
            try:
                return datetime.strptime(str(integer), "%Y%m%d").date()
            except ValueError:
                pass
        # Excel seri tarihi için makul aralık.
        if excel_epoch is not None and 1 <= number <= 100000:
            try:
                converted = from_excel(number, excel_epoch)
                return converted.date() if isinstance(converted, datetime) else converted
            except (TypeError, ValueError, OverflowError):
                pass

    text = unicodedata.normalize("NFKC", str(value)).strip()
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text)

    # Saat bilgisi varsa günlük eşleştirme için tarih kısmı kullanılır.
    candidates = [text]
    if "T" in text:
        candidates.append(text.split("T", 1)[0])
    if " " in text:
        candidates.append(text.split(" ", 1)[0])

    formats = (
        "%d.%m.%Y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
        "%d.%m.%y",
        "%d/%m/%y",
        "%d-%m-%y",
    )
    for candidate in candidates:
        for fmt in formats:
            try:
                return datetime.strptime(candidate, fmt).date()
            except ValueError:
                continue

    raise ValueError(f"Tarih olarak okunamadı: {text!r}")


def _decimal_places(number_format: str) -> int | None:
    """Basit Excel sayı biçiminden gösterilecek ondalık basamak sayısını bulur."""
    if not number_format or number_format == "General":
        return None
    first_section = number_format.split(";", 1)[0]
    first_section = re.sub(r'"[^"]*"', "", first_section)
    first_section = re.sub(r"\[[^\]]*\]", "", first_section)
    # Excel biçim kodlarında nokta ondalık, virgül binlik ayırıcıdır.
    # Bu nedenle #,##0 biçimindeki virgülü ondalık işareti sanmamak gerekir.
    match = re.search(r"\.([0#]+)", first_section)
    if match:
        return len(match.group(1))
    # Bazı elle yazılmış yerel biçimler 0,00 şeklinde gelebilir.
    local_match = re.search(r",(0+)(?:[^0#]|$)", first_section)
    return len(local_match.group(1)) if local_match else 0


def _localize_number(text: str) -> str:
    decimal_point = locale.localeconv().get("decimal_point") or "."
    if decimal_point != ".":
        return text.replace(".", decimal_point)
    return text


def format_excel_value(value: Any, number_format: str) -> str:
    """Hücre değerini Word'e yazılabilecek, okunaklı metne çevirir."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"

    if isinstance(value, Number):
        numeric = float(value)
        is_percentage = "%" in (number_format or "")
        if is_percentage:
            numeric *= 100

        places = _decimal_places(number_format)
        if places is None:
            if numeric.is_integer():
                text = str(int(numeric))
            else:
                text = format(numeric, ".15g")
        else:
            text = f"{numeric:.{places}f}"

        text = _localize_number(text)
        return f"{text}%" if is_percentage else text

    return str(value).strip()


def _header_map_from_values(values: list[Any]) -> dict[str, int] | None:
    found: dict[str, int] = {}
    for index, value in enumerate(values):
        canonical = NORMALIZED_HEADERS.get(normalize_header(value))
        if canonical is not None and canonical not in found:
            found[canonical] = index
    if all(header in found for header in REQUIRED_HEADERS):
        return found
    return None


def _find_excel_candidates(workbook: Any) -> list[tuple[int, str, int, dict[str, int]]]:
    candidates: list[tuple[int, str, int, dict[str, int]]] = []
    for worksheet in workbook.worksheets:
        max_row = min(worksheet.max_row or 0, MAX_HEADER_SCAN_ROWS)
        max_col = min(worksheet.max_column or 0, MAX_HEADER_SCAN_COLUMNS)
        for row_number in range(1, max_row + 1):
            values = [worksheet.cell(row_number, col).value for col in range(1, max_col + 1)]
            header_map = _header_map_from_values(values)
            if header_map is None:
                continue

            date_col = header_map["Date"] + 1
            valid_dates = 0
            for data_row in range(row_number + 1, worksheet.max_row + 1):
                value = worksheet.cell(data_row, date_col).value
                if value in (None, ""):
                    continue
                try:
                    parse_date(value, workbook.epoch)
                    valid_dates += 1
                except ValueError:
                    continue
            candidates.append((valid_dates, worksheet.title, row_number, header_map))
    return candidates


def read_excel(excel_path: Path) -> ExcelData:
    try:
        values_book = load_workbook(excel_path, data_only=True, read_only=False)
        formulas_book = load_workbook(excel_path, data_only=False, read_only=False)
    except Exception as exc:
        raise TransferError(f"Excel dosyası açılamadı:\n{exc}") from exc

    try:
        candidates = _find_excel_candidates(values_book)
        if not candidates:
            raise TransferError(
                "Excel içinde Date, PHF (€), Volume ve PHF VT başlıklarını aynı satırda "
                "bulamadım. İlk 100 satır ve ilk 100 sütun kontrol edildi."
            )

        candidates.sort(key=lambda item: item[0], reverse=True)
        best_count = candidates[0][0]
        best = [item for item in candidates if item[0] == best_count]
        if len(best) > 1:
            locations = ", ".join(f"{item[1]}!{item[2]}" for item in best)
            raise TransferError(
                "Excel içinde birden fazla eşit aday tablo bulundu. Başlıkları tek bir tabloda "
                f"bırakın veya diğerlerini yeniden adlandırın. Adaylar: {locations}"
            )

        _, sheet_name, header_row, header_map = best[0]
        worksheet = values_book[sheet_name]
        formula_sheet = formulas_book[sheet_name]
        records: dict[date, ExcelRecord] = {}
        duplicates: list[str] = []

        for row_number in range(header_row + 1, worksheet.max_row + 1):
            date_cell = worksheet.cell(row_number, header_map["Date"] + 1)
            if date_cell.value in (None, ""):
                continue
            try:
                record_date = parse_date(date_cell.value, values_book.epoch)
            except ValueError as exc:
                raise TransferError(
                    f"Excel'de {sheet_name}!{date_cell.coordinate} hücresindeki tarih okunamadı: "
                    f"{date_cell.value!r}"
                ) from exc

            if record_date in records:
                duplicates.append(record_date.strftime("%d.%m.%Y"))
                continue

            values: dict[str, str] = {}
            blank_headers: list[str] = []
            for header in VALUE_HEADERS:
                col_number = header_map[header] + 1
                cell = worksheet.cell(row_number, col_number)
                formula_cell = formula_sheet.cell(row_number, col_number)
                if cell.value is None and formula_cell.data_type == "f":
                    raise TransferError(
                        f"{sheet_name}!{cell.coordinate} formülünün kayıtlı sonucu boş. "
                        "Excel dosyasını Microsoft Excel'de açıp hesaplatın, kaydedin ve tekrar deneyin."
                    )
                text = format_excel_value(cell.value, cell.number_format)
                values[header] = text
                if text == "":
                    blank_headers.append(header)

            records[record_date] = ExcelRecord(values=values, blank_headers=blank_headers)

        if duplicates:
            unique = ", ".join(sorted(set(duplicates)))
            raise TransferError(f"Excel'de aynı tarih birden fazla kez bulunuyor: {unique}")
        if not records:
            raise TransferError("Excel tablosunda işlenebilecek tarih satırı bulunamadı.")

        return ExcelData(sheet_name=sheet_name, header_row=header_row, records=records)
    finally:
        values_book.close()
        formulas_book.close()


def _find_word_table(document: Any, excel_dates: set[date]) -> tuple[Any, int, int, dict[str, int]]:
    candidates: list[tuple[int, int, int, Any, dict[str, int]]] = []

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows[:MAX_WORD_HEADER_SCAN_ROWS]):
            header_map = _header_map_from_values([cell.text for cell in row.cells])
            if header_map is None:
                continue

            matched_dates = 0
            readable_dates = 0
            date_col = header_map["Date"]
            for data_row in table.rows[row_index + 1 :]:
                if date_col >= len(data_row.cells):
                    continue
                raw_date = data_row.cells[date_col].text.strip()
                if not raw_date:
                    continue
                try:
                    parsed = parse_date(raw_date)
                    readable_dates += 1
                    if parsed in excel_dates:
                        matched_dates += 1
                except ValueError:
                    continue
            candidates.append((matched_dates, readable_dates, -table_index, table, header_map | {"_row": row_index}))

    if not candidates:
        raise TransferError(
            "Word içinde Date, PHF (€), Volume ve PHF VT başlıklarını içeren tablo bulunamadı."
        )

    candidates.sort(key=lambda item: (item[0], item[1], item[2]), reverse=True)
    best = candidates[0]
    matched_dates, _, negative_table_index, table, combined_map = best
    if matched_dates == 0:
        raise TransferError(
            "Word tablosu bulundu ancak Word ve Excel arasında eşleşen tarih yok. "
            "Tarih aralıklarını kontrol edin."
        )

    table_index = -negative_table_index
    header_row = combined_map.pop("_row")
    return table, table_index, header_row, combined_map


def _set_cell_text_preserving_format(cell: Any, text: str) -> None:
    """Boş hücrenin mevcut paragraf ve yazı biçimini mümkün olduğunca korur."""
    paragraph = cell.paragraphs[0]
    first_run = None
    for current_paragraph in cell.paragraphs:
        for run in current_paragraph.runs:
            if first_run is None:
                first_run = run
            run.text = ""
    if first_run is None:
        first_run = paragraph.add_run()
    first_run.text = text


def transfer(excel_path: Path, word_path: Path, output_path: Path) -> TransferResult:
    excel_data = read_excel(excel_path)
    try:
        document = Document(word_path)
    except Exception as exc:
        raise TransferError(f"Word dosyası açılamadı:\n{exc}") from exc

    table, table_index, header_row, header_map = _find_word_table(
        document, set(excel_data.records)
    )
    result = TransferResult(
        excel_sheet=excel_data.sheet_name,
        excel_header_row=excel_data.header_row,
        word_table_number=table_index + 1,
        word_header_row=header_row + 1,
    )

    word_dates_seen: set[date] = set()
    date_col = header_map["Date"]

    for row in table.rows[header_row + 1 :]:
        if date_col >= len(row.cells):
            result.ignored_word_rows += 1
            continue
        raw_date = row.cells[date_col].text.strip()
        if not raw_date:
            continue
        try:
            row_date = parse_date(raw_date)
        except ValueError:
            result.ignored_word_rows += 1
            continue

        word_dates_seen.add(row_date)
        record = excel_data.records.get(row_date)
        if record is None:
            result.word_dates_without_excel.add(row_date)
            continue

        result.matched_dates.add(row_date)
        result.blank_excel_cells += len(record.blank_headers)
        for header in VALUE_HEADERS:
            col_index = header_map[header]
            if col_index >= len(row.cells):
                raise TransferError(
                    f"Word tablosunda {header} sütunu bazı satırlarda erişilebilir değil. "
                    "Birleştirilmiş hücreleri kontrol edin."
                )
            target_cell = row.cells[col_index]
            if target_cell.text.strip() and not OVERWRITE_EXISTING_CELLS:
                result.skipped_nonempty_cells += 1
                continue
            _set_cell_text_preserving_format(target_cell, record.values[header])

    result.excel_dates_not_in_word = set(excel_data.records) - word_dates_seen
    if not result.matched_dates:
        raise TransferError("Hiçbir tarih eşleşmedi; çıktı oluşturulmadı.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}_",
            suffix=".docx",
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        document.save(temporary_path)
        os.replace(temporary_path, output_path)
    except Exception as exc:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise TransferError(f"Çıktı Word dosyası kaydedilemedi:\n{exc}") from exc

    return result


def _short_date_list(values: set[date], limit: int = 8) -> str:
    ordered = [item.strftime("%d.%m.%Y") for item in sorted(values)]
    if len(ordered) > limit:
        return ", ".join(ordered[:limit]) + f" ve {len(ordered) - limit} tarih daha"
    return ", ".join(ordered)


def result_message(result: TransferResult, output_path: Path) -> str:
    lines = [
        "Word dosyası oluşturuldu.",
        "",
        f"Eşleşen tarih: {len(result.matched_dates)}",
        f"Excel sayfası: {result.excel_sheet}",
        f"Word tablosu: {result.word_table_number}",
        f"Çıktı: {output_path}",
    ]
    if result.word_dates_without_excel:
        lines.extend(
            [
                "",
                f"Excel'de bulunmayan Word tarihleri ({len(result.word_dates_without_excel)}): ",
                _short_date_list(result.word_dates_without_excel),
            ]
        )
    if result.excel_dates_not_in_word:
        lines.extend(
            [
                "",
                f"Word'de bulunmayan Excel tarihleri ({len(result.excel_dates_not_in_word)}): ",
                _short_date_list(result.excel_dates_not_in_word),
            ]
        )
    if result.skipped_nonempty_cells:
        lines.extend(
            [
                "",
                f"Dolu olduğu için değiştirilmemiş Word hücresi: {result.skipped_nonempty_cells}",
            ]
        )
    if result.blank_excel_cells:
        lines.extend(["", f"Excel'de boş gelen değer hücresi: {result.blank_excel_cells}"])
    if result.ignored_word_rows:
        lines.extend(["", f"Tarih olarak okunamayan Word satırı: {result.ignored_word_rows}"])
    return "\n".join(lines)


# Çok tablolı yeni aktarım düzeni. Eski tek-tablo fonksiyonları, daha önce
# oluşturulmuş entegrasyonları bozmayacak şekilde dosyada tutulmaktadır.
@dataclass(frozen=True)
class TableSpec:
    name: str
    headers: tuple[str, ...]
    date_key: bool = False


TABLE_SPECS = (
    TableSpec("PHF", ("Date", "PHF (€)", "Volume", "PHF VT"), True),
    TableSpec("HPU", ("Date", "HPU", "Volume"), True),
    TableSpec("Scrap", ("Date", "Scrap (€)", "Volume", "Scrap VT"), True),
    TableSpec("Electric Per Vehicle", ("Electric Per Vehicle (€)", "EMB", "TOL", "PEI", "MON", "CHA")),
    TableSpec("Total Gas Per Vehicle", ("Total Gas Per Vehicle (€)", "EMB", "TOL", "PEI", "MON", "CHA", "Total")),
    TableSpec("Forklift", ("Date", "Forklift Quantitiy", "Volume", "Forklift Q / Volume"), True),
    TableSpec("Ecart INV", ("Date", "Ecart INV", "Volume", "Ecart INV / Volume"), True),
)


@dataclass
class TableMapping:
    spec: TableSpec
    excel_sheet: str
    excel_header_row: int
    excel_columns: dict[str, int]
    word_table_number: int
    word_header_row: int
    word_columns: dict[str, int]
    detected: bool = True


@dataclass
class MultiResult:
    rows_by_table: dict[str, int] = field(default_factory=dict)
    unmatched_dates: dict[str, int] = field(default_factory=dict)


def _map_spec_headers(values: list[Any], spec: TableSpec) -> dict[str, int] | None:
    normalized = {normalize_header(value): index + 1 for index, value in enumerate(values) if value not in (None, "")}
    found: dict[str, int] = {}
    for header in spec.headers:
        aliases = [normalize_header(header)]
        if header == "Forklift Quantitiy":
            aliases.append(normalize_header("Forklift Quantity"))
        column = next((normalized[a] for a in aliases if a in normalized), None)
        if column is None:
            return None
        found[header] = column
    return found


def _excel_candidate_score(sheet: Any, row_number: int, columns: dict[str, int], spec: TableSpec, epoch: Any) -> int:
    score = 0
    blanks = 0
    key = "Date" if spec.date_key else spec.headers[0]
    for row in range(row_number + 1, min(sheet.max_row, row_number + 2000) + 1):
        value = sheet.cell(row, columns[key]).value
        if value in (None, ""):
            blanks += 1
            if blanks >= 5 and score:
                break
            continue
        blanks = 0
        if spec.date_key:
            try:
                parse_date(value, epoch)
            except ValueError:
                continue
        score += 1
    return score


def detect_mappings(excel_path: Path, word_path: Path) -> tuple[list[TableMapping], list[str]]:
    try:
        workbook = load_workbook(excel_path, data_only=True, read_only=False)
        document = Document(word_path)
    except Exception as exc:
        raise TransferError(f"Dosyalar otomatik taranamadı:\n{exc}") from exc

    mappings: list[TableMapping] = []
    try:
        sheet_names = workbook.sheetnames
        for spec_index, spec in enumerate(TABLE_SPECS):
            excel_candidates = []
            for sheet in workbook.worksheets:
                for row_number in range(1, min(sheet.max_row, MAX_HEADER_SCAN_ROWS) + 1):
                    values = [sheet.cell(row_number, col).value for col in range(1, min(sheet.max_column, 200) + 1)]
                    columns = _map_spec_headers(values, spec)
                    if columns:
                        score = _excel_candidate_score(sheet, row_number, columns, spec, workbook.epoch)
                        excel_candidates.append((score, sheet.title, row_number, columns))
            excel_candidates.sort(key=lambda item: item[0], reverse=True)

            word_candidates = []
            for table_number, table in enumerate(document.tables, 1):
                for row_number, row in enumerate(table.rows[:MAX_WORD_HEADER_SCAN_ROWS], 1):
                    columns = _map_spec_headers([cell.text for cell in row.cells], spec)
                    if columns:
                        word_candidates.append((table_number, row_number, columns))

            detected = bool(excel_candidates and word_candidates)
            if excel_candidates:
                _, sheet_name, excel_row, excel_columns = excel_candidates[0]
            else:
                sheet_name, excel_row = sheet_names[0], 1
                excel_columns = {header: i + 1 for i, header in enumerate(spec.headers)}
            if word_candidates:
                word_table, word_row, word_columns = word_candidates[0]
            else:
                word_table, word_row = min(spec_index + 1, max(len(document.tables), 1)), 1
                word_columns = {header: i + 1 for i, header in enumerate(spec.headers)}
            mappings.append(TableMapping(spec, sheet_name, excel_row, excel_columns, word_table, word_row, word_columns, detected))
        return mappings, sheet_names
    finally:
        workbook.close()


class MappingDialog(tk.Toplevel):
    """Otomatik seçimleri aktarım öncesinde zorunlu olarak kullanıcıya sorar."""

    def __init__(self, master: tk.Widget, mappings: list[TableMapping], sheet_names: list[str]):
        super().__init__(master)
        self.title("Tablo ve sütun eşleştirmelerini kontrol edin")
        self.geometry("860x580")
        self.transient(master)
        self.grab_set()
        self.result: list[TableMapping] | None = None
        self._mappings = mappings
        self._vars = []

        ttk.Label(self, text="Otomatik seçimleri kontrol edin. Gerekirse değiştirip Onayla ve Aktar düğmesine basın.", wraplength=820).pack(anchor="w", padx=12, pady=(12, 6))
        notebook = ttk.Notebook(self)
        notebook.pack(fill="both", expand=True, padx=12, pady=6)

        for mapping in mappings:
            frame = ttk.Frame(notebook, padding=12)
            notebook.add(frame, text=mapping.spec.name)
            status = "Otomatik bulundu" if mapping.detected else "Otomatik bulunamadı - elle kontrol edin"
            ttk.Label(frame, text=status).grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 10))
            sheet_var = tk.StringVar(value=mapping.excel_sheet)
            excel_row_var = tk.StringVar(value=str(mapping.excel_header_row))
            word_table_var = tk.StringVar(value=str(mapping.word_table_number))
            word_row_var = tk.StringVar(value=str(mapping.word_header_row))
            ttk.Label(frame, text="Excel sayfası").grid(row=1, column=0, sticky="w")
            ttk.Combobox(frame, textvariable=sheet_var, values=sheet_names, state="readonly", width=28).grid(row=1, column=1, sticky="w")
            ttk.Label(frame, text="Excel başlık satırı").grid(row=1, column=2, sticky="e", padx=(20, 5))
            ttk.Entry(frame, textvariable=excel_row_var, width=8).grid(row=1, column=3, sticky="w")
            ttk.Label(frame, text="Word tablo numarası").grid(row=2, column=0, sticky="w", pady=6)
            ttk.Entry(frame, textvariable=word_table_var, width=8).grid(row=2, column=1, sticky="w", pady=6)
            ttk.Label(frame, text="Word başlık satırı").grid(row=2, column=2, sticky="e", padx=(20, 5), pady=6)
            ttk.Entry(frame, textvariable=word_row_var, width=8).grid(row=2, column=3, sticky="w", pady=6)
            ttk.Label(frame, text="Alan").grid(row=3, column=0, sticky="w", pady=(10, 3))
            ttk.Label(frame, text="Excel sütun no").grid(row=3, column=1, sticky="w", pady=(10, 3))
            ttk.Label(frame, text="Word sütun no").grid(row=3, column=2, sticky="w", pady=(10, 3))
            column_vars = {}
            for i, header in enumerate(mapping.spec.headers, 4):
                ev = tk.StringVar(value=str(mapping.excel_columns[header]))
                wv = tk.StringVar(value=str(mapping.word_columns[header]))
                ttk.Label(frame, text=header).grid(row=i, column=0, sticky="w", pady=2)
                ttk.Entry(frame, textvariable=ev, width=10).grid(row=i, column=1, sticky="w", pady=2)
                ttk.Entry(frame, textvariable=wv, width=10).grid(row=i, column=2, sticky="w", pady=2)
                column_vars[header] = (ev, wv)
            self._vars.append((mapping, sheet_var, excel_row_var, word_table_var, word_row_var, column_vars))

        buttons = ttk.Frame(self)
        buttons.pack(fill="x", padx=12, pady=(4, 12))
        ttk.Button(buttons, text="İptal", command=self.destroy).pack(side="right")
        ttk.Button(buttons, text="Onayla ve Aktar", command=self._confirm).pack(side="right", padx=(0, 8))
        self.protocol("WM_DELETE_WINDOW", self.destroy)

    def _confirm(self):
        try:
            result = []
            for mapping, sheet, erow, wtable, wrow, column_vars in self._vars:
                excel_columns = {h: int(v[0].get()) for h, v in column_vars.items()}
                word_columns = {h: int(v[1].get()) for h, v in column_vars.items()}
                values = [int(erow.get()), int(wtable.get()), int(wrow.get()), *excel_columns.values(), *word_columns.values()]
                if any(value < 1 for value in values):
                    raise ValueError
                result.append(TableMapping(mapping.spec, sheet.get(), int(erow.get()), excel_columns, int(wtable.get()), int(wrow.get()), word_columns, mapping.detected))
        except ValueError:
            messagebox.showerror("Geçersiz eşleştirme", "Satır, tablo ve sütun numaraları 1 veya daha büyük tam sayı olmalıdır.", parent=self)
            return
        self.result = result
        self.destroy()


def _excel_rows(sheet: Any, mapping: TableMapping, epoch: Any) -> list[tuple[date | None, dict[str, str]]]:
    rows = []
    started = False
    blanks = 0
    key_header = "Date" if mapping.spec.date_key else mapping.spec.headers[0]
    for row_number in range(mapping.excel_header_row + 1, sheet.max_row + 1):
        scan_values = [sheet.cell(row_number, col).value for col in range(1, min(sheet.max_column, 200) + 1)]
        if started and any(_map_spec_headers(scan_values, other_spec) for other_spec in TABLE_SPECS):
            break
        key_cell = sheet.cell(row_number, mapping.excel_columns[key_header])
        key_value = key_cell.value
        row_date = None
        if mapping.spec.date_key:
            try:
                row_date = parse_date(key_value, epoch)
            except ValueError:
                if started and key_value in (None, ""):
                    blanks += 1
                    if blanks >= 5:
                        break
                continue
        elif key_value in (None, ""):
            if started:
                blanks += 1
                if blanks >= 5:
                    break
            continue
        started, blanks = True, 0
        values = {}
        for header in mapping.spec.headers:
            cell = sheet.cell(row_number, mapping.excel_columns[header])
            values[header] = format_excel_value(cell.value, cell.number_format)
        rows.append((row_date, values))
    if not rows:
        raise TransferError(f"{mapping.spec.name}: Excel'de başlangıç verisi bulunamadı.")
    return rows


def transfer_multi(excel_path: Path, word_path: Path, output_path: Path, mappings: list[TableMapping]) -> MultiResult:
    try:
        workbook = load_workbook(excel_path, data_only=True, read_only=False)
        document = Document(word_path)
    except Exception as exc:
        raise TransferError(f"Dosyalar açılamadı:\n{exc}") from exc
    result = MultiResult()
    try:
        for mapping in mappings:
            if mapping.excel_sheet not in workbook.sheetnames:
                raise TransferError(f"{mapping.spec.name}: Excel sayfası bulunamadı: {mapping.excel_sheet}")
            if mapping.word_table_number > len(document.tables):
                raise TransferError(f"{mapping.spec.name}: Word tablo numarası mevcut değil: {mapping.word_table_number}")
            sheet = workbook[mapping.excel_sheet]
            table = document.tables[mapping.word_table_number - 1]
            rows = _excel_rows(sheet, mapping, workbook.epoch)
            copied = 0
            unmatched = 0

            if mapping.spec.date_key:
                dated_word_rows = {}
                for row in table.rows[mapping.word_header_row:]:
                    col = mapping.word_columns["Date"] - 1
                    if col >= len(row.cells):
                        continue
                    try:
                        dated_word_rows[parse_date(row.cells[col].text)] = row
                    except ValueError:
                        continue
                if dated_word_rows:
                    for row_date, values in rows:
                        target = dated_word_rows.get(row_date)
                        if target is None:
                            unmatched += 1
                            continue
                        for header in mapping.spec.headers:
                            if header != "Date":
                                _set_cell_text_preserving_format(target.cells[mapping.word_columns[header] - 1], values[header])
                        copied += 1
                else:
                    while len(table.rows) - mapping.word_header_row < len(rows):
                        table.add_row()
                    for target, (_, values) in zip(table.rows[mapping.word_header_row:], rows):
                        for header in mapping.spec.headers:
                            _set_cell_text_preserving_format(target.cells[mapping.word_columns[header] - 1], values[header])
                        copied += 1
            else:
                while len(table.rows) - mapping.word_header_row < len(rows):
                    table.add_row()
                for target, (_, values) in zip(table.rows[mapping.word_header_row:], rows):
                    for header in mapping.spec.headers:
                        _set_cell_text_preserving_format(target.cells[mapping.word_columns[header] - 1], values[header])
                    copied += 1
            result.rows_by_table[mapping.spec.name] = copied
            result.unmatched_dates[mapping.spec.name] = unmatched
    finally:
        workbook.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(f".{output_path.stem}_gecici.docx")
    try:
        document.save(temp_path)
        os.replace(temp_path, output_path)
    except Exception as exc:
        temp_path.unlink(missing_ok=True)
        raise TransferError(f"Çıktı kaydedilemedi:\n{exc}") from exc
    return result


def multi_result_message(result: MultiResult, output_path: Path) -> str:
    lines = ["Word dosyası oluşturuldu.", ""]
    for spec in TABLE_SPECS:
        copied = result.rows_by_table.get(spec.name, 0)
        unmatched = result.unmatched_dates.get(spec.name, 0)
        suffix = f"; Word'de tarihi bulunamayan: {unmatched}" if unmatched else ""
        lines.append(f"{spec.name}: {copied} satır aktarıldı{suffix}")
    lines.extend(["", f"Çıktı: {output_path}"])
    return "\n".join(lines)


class Application(ttk.Frame):
    def __init__(self, master: tk.Tk) -> None:
        super().__init__(master, padding=16)
        self.grid(sticky="nsew")
        master.columnconfigure(0, weight=1)
        master.rowconfigure(0, weight=1)
        self.columnconfigure(1, weight=1)

        self.excel_var = tk.StringVar()
        self.word_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.status_var = tk.StringVar(value="Dosyaları seçin.")

        self._file_row(0, "Excel dosyası", self.excel_var, self.choose_excel)
        self._file_row(1, "Word şablonu", self.word_var, self.choose_word)
        self._file_row(2, "Çıktı Word", self.output_var, self.choose_output)

        ttk.Separator(self, orient="horizontal").grid(
            row=3, column=0, columnspan=3, sticky="ew", pady=(14, 12)
        )
        self.run_button = ttk.Button(self, text="Eşleştirmeleri Kontrol Et ve Oluştur", command=self.run_transfer)
        self.run_button.grid(row=4, column=0, columnspan=3, sticky="ew")
        ttk.Label(self, textvariable=self.status_var, wraplength=640).grid(
            row=5, column=0, columnspan=3, sticky="w", pady=(12, 0)
        )

    def _file_row(self, row: int, label: str, variable: tk.StringVar, command: Any) -> None:
        ttk.Label(self, text=label).grid(row=row, column=0, sticky="w", padx=(0, 10), pady=5)
        ttk.Entry(self, textvariable=variable).grid(row=row, column=1, sticky="ew", pady=5)
        ttk.Button(self, text="Seç", command=command).grid(row=row, column=2, padx=(10, 0), pady=5)

    def choose_excel(self) -> None:
        path = filedialog.askopenfilename(
            title="Excel dosyasını seçin",
            filetypes=[("Excel dosyaları", "*.xlsx *.xlsm"), ("Tüm dosyalar", "*.*")],
        )
        if path:
            self.excel_var.set(path)

    def choose_word(self) -> None:
        path = filedialog.askopenfilename(
            title="Word şablonunu seçin",
            filetypes=[("Word dosyaları", "*.docx"), ("Tüm dosyalar", "*.*")],
        )
        if path:
            self.word_var.set(path)
            word_path = Path(path)
            default_output = word_path.with_name(f"{word_path.stem}_doldurulmus.docx")
            self.output_var.set(str(default_output))

    def choose_output(self) -> None:
        initial = Path(self.output_var.get()) if self.output_var.get() else Path.cwd() / "sonuc.docx"
        path = filedialog.asksaveasfilename(
            title="Çıktı Word dosyasını kaydedin",
            defaultextension=".docx",
            initialdir=str(initial.parent),
            initialfile=initial.name,
            filetypes=[("Word dosyası", "*.docx")],
        )
        if path:
            self.output_var.set(path)

    def run_transfer(self) -> None:
        excel_path = Path(self.excel_var.get().strip())
        word_path = Path(self.word_var.get().strip())
        output_path = Path(self.output_var.get().strip())

        if not self.excel_var.get().strip() or not excel_path.is_file():
            messagebox.showerror("Eksik dosya", "Geçerli bir Excel dosyası seçin.")
            return
        if not self.word_var.get().strip() or not word_path.is_file():
            messagebox.showerror("Eksik dosya", "Geçerli bir Word şablonu seçin.")
            return
        if not self.output_var.get().strip():
            messagebox.showerror("Eksik çıktı", "Çıktı Word dosyasının yolunu seçin.")
            return
        if output_path.suffix.lower() != ".docx":
            output_path = output_path.with_suffix(".docx")
            self.output_var.set(str(output_path))
        if output_path.resolve() == word_path.resolve():
            messagebox.showerror("Geçersiz çıktı", "Çıktı yolu Word şablonuyla aynı olamaz.")
            return
        if output_path.exists():
            overwrite = messagebox.askyesno(
                "Dosya mevcut", f"Bu çıktı dosyası zaten var. Üzerine yazılsın mı?\n\n{output_path}"
            )
            if not overwrite:
                return

        self.run_button.state(["disabled"])
        self.status_var.set("Tablolar ve sütunlar otomatik aranıyor...")
        self.update_idletasks()
        try:
            mappings, sheet_names = detect_mappings(excel_path, word_path)
            dialog = MappingDialog(self, mappings, sheet_names)
            self.wait_window(dialog)
            if dialog.result is None:
                self.status_var.set("İşlem iptal edildi.")
                return
            self.status_var.set("Veriler Word dosyasına aktarılıyor...")
            self.update_idletasks()
            result = transfer_multi(excel_path, word_path, output_path, dialog.result)
        except TransferError as exc:
            self.status_var.set("İşlem tamamlanamadı.")
            messagebox.showerror("İşlem tamamlanamadı", str(exc))
        except Exception as exc:
            self.status_var.set("Beklenmeyen bir hata oluştu.")
            messagebox.showerror(
                "Beklenmeyen hata",
                f"İşlem sırasında beklenmeyen bir hata oluştu:\n{exc}",
            )
        else:
            total_rows = sum(result.rows_by_table.values())
            self.status_var.set(f"Tamamlandı: {total_rows} tablo satırı aktarıldı.")
            messagebox.showinfo("Tamamlandı", multi_result_message(result, output_path))
        finally:
            self.run_button.state(["!disabled"])


def main() -> None:
    try:
        locale.setlocale(locale.LC_ALL, "")
    except locale.Error:
        pass

    root = tk.Tk()
    root.title("Excel Word Aktarıcı")
    root.minsize(720, 245)
    Application(root)
    root.mainloop()


if __name__ == "__main__":
    main()
